"""Unit tests for BRDGenerator, DOCX rendering, and BRD document version persistence."""

from datetime import datetime, timezone
import io
from unittest.mock import MagicMock
import uuid
import pytest
from docx import Document

from app.core.brd_generator import BRDGenerator
from app.core.docx_renderer import render_brd_to_docx
from app.db.repository import ProjectRepository
from app.models.assumption import Assumption, AssumptionStatus, ImpactLevel
from app.models.document import (
    BRDDocumentVersion,
    SourceDocument,
    SourceDocumentStatus,
)
from app.models.project import Project, ProjectStatus
from app.models.requirement import Priority, Requirement, RequirementType
from tests.test_api_persistence import InMemoryProjectRepository


def _create_sample_project_and_data():
    """Helper creating a sample project, grounded requirements, and assumptions."""
    project_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc)
    project = Project(
        id=project_id,
        title="Fleet Telematics Management Platform",
        idea="A comprehensive fleet management system tracking vehicle telemetry and route dispatching.",
        status=ProjectStatus.READY_TO_GENERATE,
        input_revision=2,
        created_at=now,
        updated_at=now,
    )

    doc_id = str(uuid.uuid4())
    source_doc = SourceDocument(
        id=doc_id,
        project_id=project_id,
        filename="fleet_specs.pdf",
        file_type="pdf",
        status=SourceDocumentStatus.READY,
    )

    requirements = [
        Requirement(
            id="REQ-BIZ-001",
            title="Reduce Fuel Consumption",
            statement="The platform shall optimize dispatch routes to decrease overall fleet fuel burn by 15%.",
            type=RequirementType.BUSINESS,
            priority=Priority.MUST,
            actor="Fleet Operations Executive",
            outcome="Measurable reduction in operational expenditures.",
            acceptance_criteria=["Fuel expenditure reports show >= 15% reduction quarterly."],
            trace_refs=["ANSWER:A-101"],
            source_chunk_ids=[str(uuid.uuid4())],
        ),
        Requirement(
            id="REQ-FUNC-001",
            title="Real-Time GPS Vehicle Tracking",
            statement="The system shall capture and broadcast GPS coordinates every 10 seconds per vehicle.",
            type=RequirementType.FUNCTIONAL,
            priority=Priority.MUST,
            actor="Dispatch Operator",
            outcome="Real-time situational visibility on map interface.",
            acceptance_criteria=[
                "GPS ping latency is under 3 seconds.",
                "Vehicle icon updates dynamically on dispatcher dashboard.",
            ],
            trace_refs=["CHUNK:fleet_specs-p2"],
            source_chunk_ids=[str(uuid.uuid4())],
        ),
        Requirement(
            id="REQ-NFR-001",
            title="API Latency & Throughput",
            statement="Telemetry ingestion API shall process 5,000 requests per second with p99 latency < 200ms.",
            type=RequirementType.NONFUNCTIONAL,
            priority=Priority.MUST,
            actor="Telemetry Ingestion Gateway",
            outcome="Zero message dropping during peak operational spikes.",
            acceptance_criteria=["Load tests sustain 5k req/s for 1 hour with < 200ms p99 latency."],
            trace_refs=["CHUNK:fleet_specs-p5"],
            source_chunk_ids=[str(uuid.uuid4())],
        ),
        Requirement(
            id="REQ-CONST-001",
            title="PostgreSQL Cloud Relational Storage",
            statement="All transactional dispatch logs and location historical records must be persisted in PostgreSQL.",
            type=RequirementType.CONSTRAINT,
            priority=Priority.SHOULD,
            actor="System Architect",
            outcome="Compliance with enterprise relational data governance policies.",
            acceptance_criteria=["Transactional dispatch schemas conform to 3NF standards."],
            trace_refs=["ASSUMPTION:ASM-001"],
            source_chunk_ids=[],
        ),
    ]

    assumptions = [
        Assumption(
            id="ASM-001",
            statement="Vehicles are equipped with OBD-II compliant cellular telemetry transmitters.",
            rationale="Standard industry practice for commercial vehicle fleets.",
            impact=ImpactLevel.HIGH,
            status=AssumptionStatus.ACCEPTED,
            linked_requirement_ids=["REQ-FUNC-001"],
        ),
        Assumption(
            id="ASM-002",
            statement="Cellular connectivity is available across 98% of standard delivery routes.",
            rationale="Transmitter caches up to 2 hours of data offline during dead zones.",
            impact=ImpactLevel.MEDIUM,
            status=AssumptionStatus.ACCEPTED,
            linked_requirement_ids=["REQ-FUNC-001"],
        ),
    ]

    return project, requirements, assumptions, [source_doc]


def test_brd_generator_compiles_complete_json():
    """Verify BRDGenerator compiles complete 7-section structured JSON conforming to specification."""
    project, reqs, asms, docs = _create_sample_project_and_data()
    generator = BRDGenerator()

    brd_version = generator.compile_brd(
        project=project,
        requirements=reqs,
        assumptions=asms,
        source_documents=docs,
        version_number=1,
        strict_quality=True,
    )

    # Top-level model validations
    assert brd_version.project_id == project.id
    assert brd_version.version == 1
    assert brd_version.version_number == 1
    assert brd_version.title == project.title
    assert brd_version.quality_check_passed is True

    content = brd_version.content_json
    assert content["title"] == project.title
    assert content["version"] == 1
    assert content["quality_check_passed"] is True

    # 1. Executive Summary & Objectives
    assert "executive_summary" in content
    assert len(content["executive_summary"]) > 50
    assert "project_objectives" in content
    assert len(content["project_objectives"]) >= 1

    # 2. Scope (In-Scope & Out-of-Scope)
    assert "scope" in content
    assert len(content["scope"]["in_scope"]) >= 1
    assert len(content["scope"]["out_of_scope"]) >= 1

    # 3. Stakeholders & User Personas
    assert "stakeholders_and_personas" in content
    personas = content["stakeholders_and_personas"]
    assert len(personas) >= 1
    assert any(p["persona"] == "Dispatch Operator" for p in personas)

    # 4. Functional Requirements
    assert "functional_requirements" in content
    func_reqs = content["functional_requirements"]
    assert len(func_reqs) == 1
    assert func_reqs[0]["id"] == "REQ-FUNC-001"
    assert func_reqs[0]["actor"] == "Dispatch Operator"
    assert len(func_reqs[0]["acceptance_criteria"]) == 2
    assert len(func_reqs[0]["trace_refs"]) > 0

    # 5. Non-Functional Requirements & Constraints
    assert "nonfunctional_requirements" in content
    assert len(content["nonfunctional_requirements"]) == 1
    assert content["nonfunctional_requirements"][0]["id"] == "REQ-NFR-001"
    assert "constraints" in content
    assert len(content["constraints"]) == 1
    assert content["constraints"][0]["id"] == "REQ-CONST-001"

    # 6. Assumptions & Dependencies
    assert "assumptions" in content
    assert len(content["assumptions"]) == 2
    assert content["assumptions"][0]["id"] == "ASM-001"
    assert "dependencies" in content
    assert len(content["dependencies"]) >= 1

    # 7. Traceability Matrix summary
    assert "traceability_matrix" in content
    matrix = content["traceability_matrix"]
    assert len(matrix) == 4
    for entry in matrix:
        assert entry["grounding_status"] == "GROUNDED"
        assert len(entry["sources"]) > 0


def test_brd_generator_quality_check_detects_ungrounded():
    """Verify deterministic quality check catches ungrounded requirements and strict_quality raises."""
    project, reqs, asms, docs = _create_sample_project_and_data()

    # Create ungrounded requirement (no trace_refs and no source_chunk_ids)
    ungrounded_req = Requirement(
        id="REQ-UNGROUNDED-999",
        title="Predictive AI Maintenance",
        statement="The system shall predict alternator failure using deep learning neural networks.",
        type=RequirementType.FUNCTIONAL,
        priority=Priority.COULD,
        actor="Maintenance Technician",
        outcome="Early warning alerts.",
        acceptance_criteria=["Predictive accuracy exceeds 90%."],
        trace_refs=[],
        source_chunk_ids=[],
    )
    reqs.append(ungrounded_req)

    generator = BRDGenerator()
    passed, ungrounded = generator.check_traceability_quality(reqs)
    assert passed is False
    assert "REQ-UNGROUNDED-999" in ungrounded

    # When strict_quality is True, should raise ValueError
    with pytest.raises(ValueError, match="BRD quality check failed.*REQ-UNGROUNDED-999"):
        generator.compile_brd(
            project=project,
            requirements=reqs,
            assumptions=asms,
            source_documents=docs,
            strict_quality=True,
        )

    # When strict_quality is False, should compile but flag quality_check_passed=False
    brd = generator.compile_brd(
        project=project,
        requirements=reqs,
        assumptions=asms,
        source_documents=docs,
        strict_quality=False,
    )
    assert brd.quality_check_passed is False
    assert brd.content_json["quality_check_passed"] is False
    assert brd.content_json["quality_report"]["ungrounded_count"] == 1


def test_docx_renderer_produces_valid_docx():
    """Verify render_brd_to_docx outputs valid Microsoft Word document bytes readable by python-docx."""
    project, reqs, asms, docs = _create_sample_project_and_data()
    generator = BRDGenerator()
    brd_version = generator.compile_brd(project=project, requirements=reqs, assumptions=asms, source_documents=docs)

    output_stream = io.BytesIO()
    docx_bytes = render_brd_to_docx(brd_version.content_json, output_stream=output_stream)

    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000
    # DOCX files are zip archives starting with PK\x03\x04
    assert docx_bytes[:4] == b"PK\x03\x04"

    # Load back with docx.Document to verify structure and contents
    doc = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Fleet Telematics Management Platform" in all_text
    assert "1. Executive Summary & Project Objectives" in all_text
    assert "2. Project Scope" in all_text
    assert "3. Stakeholders & User Personas" in all_text
    assert "4. Functional Requirements" in all_text
    assert "5. Non-Functional Requirements & Constraints" in all_text
    assert "6. Assumptions & Dependencies" in all_text
    assert "7. Traceability Matrix Summary" in all_text

    # Verify tables were rendered
    assert len(doc.tables) >= 4  # Metadata, Personas, Functional Reqs, NFRs, Traceability


def test_docx_renderer_fallback_on_minimal_content():
    """Verify render_brd_to_docx handles minimal/empty dictionaries gracefully without crashing."""
    minimal_content = {
        "title": "Minimal Platform",
        "version": 1,
    }
    docx_bytes = render_brd_to_docx(minimal_content)
    assert isinstance(docx_bytes, bytes)
    assert docx_bytes[:4] == b"PK\x03\x04"

    doc = Document(io.BytesIO(docx_bytes))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Minimal Platform" in all_text


def test_project_repository_brd_version_crud():
    """Verify ProjectRepository saves, retrieves, and fetches latest BRDDocumentVersion with Supabase."""
    mock_client = MagicMock()
    repo = ProjectRepository(mock_client)

    proj_id = str(uuid.uuid4())
    version_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc)

    brd = BRDDocumentVersion(
        id=version_id,
        project_id=proj_id,
        version=1,
        input_revision=1,
        title="Fleet Telematics BRD",
        content_json={"title": "Fleet Telematics BRD", "sections": []},
        docx_path="documents/brd_v1.docx",
        model_used="llama-3.3-70b-versatile",
        prompt_version="v1.0",
        generated_at=now,
        quality_check_passed=True,
        is_current=True,
    )

    row_data = {
        "id": version_id,
        "project_id": proj_id,
        "version": 1,
        "input_revision": 1,
        "content_json": {"title": "Fleet Telematics BRD", "sections": []},
        "docx_path": "documents/brd_v1.docx",
        "model_used": "llama-3.3-70b-versatile",
        "prompt_version": "v1.0",
        "generated_at": now.isoformat(),
        "quality_check_passed": True,
        "is_current": True,
        "status": "COMPLETED",
    }

    mock_client.table().upsert().execute.return_value = MagicMock(data=[row_data])
    mock_client.table().select().eq().order().execute.return_value = MagicMock(data=[row_data])
    mock_client.table().select().eq().order().limit().execute.return_value = MagicMock(data=[row_data])

    # 1. Test save_brd_version
    saved = repo.save_brd_version(brd)
    assert saved.id == version_id
    assert saved.project_id == proj_id
    assert saved.version == 1
    assert saved.title == "Fleet Telematics BRD"

    # 2. Test get_brd_versions_by_project
    versions = repo.get_brd_versions_by_project(proj_id)
    assert len(versions) == 1
    assert versions[0].id == version_id

    # 3. Test get_latest_brd_version
    latest = repo.get_latest_brd_version(proj_id)
    assert latest is not None
    assert latest.id == version_id


def test_in_memory_repository_brd_version_crud():
    """Verify InMemoryProjectRepository supports BRD version operations seamlessly."""
    repo = InMemoryProjectRepository()
    proj_id = str(uuid.uuid4())

    brd1 = BRDDocumentVersion(
        id=str(uuid.uuid4()),
        project_id=proj_id,
        version=1,
        content_json={"title": "Fleet BRD v1"},
    )
    brd2 = BRDDocumentVersion(
        id=str(uuid.uuid4()),
        project_id=proj_id,
        version=2,
        content_json={"title": "Fleet BRD v2"},
    )

    repo.save_brd_version(brd1)
    repo.save_brd_version(brd2)

    versions = repo.get_brd_versions_by_project(proj_id)
    assert len(versions) == 2
    assert [v.version for v in versions] == [1, 2]

    latest = repo.get_latest_brd_version(proj_id)
    assert latest is not None
    assert latest.version == 2
    assert latest.content_json["title"] == "Fleet BRD v2"
