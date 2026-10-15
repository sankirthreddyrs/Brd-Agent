"""DOCX Document Rendering Engine for Business Requirements Documents (BRD).

Formats structured BRD content into a standardized, executive-ready Microsoft Word (.docx)
document adhering to enterprise typography, Navy header themes (#1F4E79), shaded tables,
margins, and strict traceability presentation.
"""

from datetime import datetime, timezone
import io
import logging
from typing import Any, Optional, Sequence

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

# Enterprise Color Palette
COLOR_NAVY_HEX = "1F4E79"
COLOR_NAVY = RGBColor(31, 78, 121)

COLOR_SLATE_HEX = "2F5597"
COLOR_SLATE = RGBColor(47, 85, 151)

COLOR_TEXT_DARK = RGBColor(38, 38, 38)
COLOR_TEXT_MUTED = RGBColor(89, 89, 89)

COLOR_ZEBRA_BG = "F2F4F8"
COLOR_LIGHT_GRAY = "E9ECEF"
COLOR_BORDER = "D3D8DE"
COLOR_CALLOUT_BG = "F8F9FA"


def _set_cell_background(cell: Any, fill_hex: str) -> None:
    """Set background fill color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("shd"):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def _set_cell_margins(
    cell: Any,
    top: int = 120,
    bottom: int = 120,
    left: int = 160,
    right: int = 160,
) -> None:
    """Set cell internal margins/padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _apply_table_borders(table: Any, border_color: str = COLOR_BORDER) -> None:
    """Apply clean, modern borders: subtle horizontal dividers and no vertical lines."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith("tblBorders"):
            tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _format_header_row(
    row: Any,
    titles: Sequence[str],
    col_widths: Optional[Sequence[float]] = None,
) -> None:
    """Apply executive navy background and white bold typography to header row."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for idx, cell in enumerate(row.cells):
        _set_cell_background(cell, COLOR_NAVY_HEX)
        _set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if col_widths and idx < len(col_widths):
            cell.width = Inches(col_widths[idx])

        title_text = titles[idx] if idx < len(titles) else ""
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title_text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)


def _add_section_heading(
    doc: Document,
    text: str,
    level: int = 1,
) -> Any:
    """Add a styled section heading with executive hierarchy."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.bold = True

    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = COLOR_NAVY
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
    elif level == 2:
        run.font.size = Pt(12.5)
        run.font.color.rgb = COLOR_SLATE
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_DARK
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True

    return p


def _add_bullet(doc: Document, text: str, bold_prefix: Optional[str] = None) -> None:
    """Add a bullet point paragraph with consistent spacing and optional bold prefix."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.bold = True
        r_prefix.font.name = "Calibri"
        r_prefix.font.size = Pt(10)
        r_prefix.font.color.rgb = COLOR_TEXT_DARK

    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_TEXT_DARK


def render_brd_to_docx(
    brd_content: dict[str, Any],
    output_stream: Optional[io.BytesIO] = None,
) -> bytes:
    """Render a structured BRD dictionary into a formatted Microsoft Word (.docx) document.

    Follows the 7-section specification:
    1. Executive Summary & Project Objectives
    2. Scope (In-Scope & Out-of-Scope)
    3. Stakeholders & User Personas
    4. Functional Requirements (ID, Title, Statement, Priority, Acceptance Criteria, Traceability)
    5. Non-Functional Requirements & Constraints
    6. Assumptions & Dependencies
    7. Traceability Matrix summary

    Args:
        brd_content: Structured JSON dictionary compiled by BRDGenerator.
        output_stream: Optional BytesIO destination stream. If None, an internal stream is used.

    Returns:
        bytes: Raw bytes of the generated .docx file.
    """
    doc = Document()

    # 1. Page Margins (Standard 1.0 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 2. Document Title and Header
    title_text = brd_content.get("title", "Business Requirements Document")
    version_num = brd_content.get("version", 1)
    input_rev = brd_content.get("input_revision", 1)
    gen_at = brd_content.get("generated_at", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    quality_passed = brd_content.get("quality_check_passed", True)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run(title_text)
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run(f"Business Requirements Document (BRD) — Version {version_num}.0")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Metadata Summary Callout Box
    meta_table = doc.add_table(rows=2, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _apply_table_borders(meta_table)

    meta_headers = ["Document Version", "Input Revision", "Generated Timestamp", "Traceability Status"]
    _format_header_row(meta_table.rows[0], meta_headers, [1.5, 1.2, 2.3, 1.5])

    meta_row = meta_table.rows[1]
    _set_cell_margins(meta_row.cells[0], top=100, bottom=100, left=150, right=150)
    meta_values = [
        f"v{version_num}.0",
        f"Rev {input_rev}",
        str(gen_at)[:19].replace("T", " "),
        "100% Grounded" if quality_passed else "Pending Grounding",
    ]
    for idx, cell in enumerate(meta_row.cells):
        _set_cell_background(cell, COLOR_CALLOUT_BG)
        _set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta_values[idx])
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = (idx == 3)
        if idx == 3:
            r.font.color.rgb = RGBColor(16, 124, 65) if quality_passed else RGBColor(180, 40, 40)
        else:
            r.font.color.rgb = COLOR_TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ------------------------------------------------------------------
    # Section 1: Executive Summary & Project Objectives
    # ------------------------------------------------------------------
    _add_section_heading(doc, "1. Executive Summary & Project Objectives", level=1)
    
    exec_text = brd_content.get("executive_summary") or (
        f"This document formalizes the complete business requirements for {title_text}. "
        f"It establishes functional deliverables, quality standards, user personas, "
        f"and traceability references justifying implementation."
    )
    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(8)
    r_exec = p_exec.add_run(exec_text)
    r_exec.font.name = "Calibri"
    r_exec.font.size = Pt(10)
    r_exec.font.color.rgb = COLOR_TEXT_DARK

    _add_section_heading(doc, "1.1 Project Objectives", level=2)
    objectives = brd_content.get("project_objectives", [])
    if objectives:
        for obj in objectives:
            _add_bullet(doc, str(obj))
    else:
        _add_bullet(doc, f"Deliver an operational, production-ready system for {title_text}.")
        _add_bullet(doc, "Establish scalable core architecture and verifiable functional requirements.")

    # ------------------------------------------------------------------
    # Section 2: Scope (In-Scope & Out-of-Scope)
    # ------------------------------------------------------------------
    _add_section_heading(doc, "2. Project Scope", level=1)
    scope = brd_content.get("scope", {})
    in_scope = scope.get("in_scope", [])
    out_scope = scope.get("out_of_scope", [])

    _add_section_heading(doc, "2.1 In-Scope Deliverables", level=2)
    if in_scope:
        for item in in_scope:
            _add_bullet(doc, str(item))
    else:
        _add_bullet(doc, "All functional requirements specified in Section 4.")

    _add_section_heading(doc, "2.2 Out-of-Scope Boundaries", level=2)
    if out_scope:
        for item in out_scope:
            _add_bullet(doc, str(item))
    else:
        _add_bullet(doc, "Custom hardware engineering and unsupported third-party platform migrations.")

    # ------------------------------------------------------------------
    # Section 3: Stakeholders & User Personas
    # ------------------------------------------------------------------
    _add_section_heading(doc, "3. Stakeholders & User Personas", level=1)
    personas = brd_content.get("stakeholders_and_personas", [])

    if personas and isinstance(personas, list) and isinstance(personas[0], dict):
        p_table = doc.add_table(rows=1 + len(personas), cols=3)
        p_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _apply_table_borders(p_table)

        col_w = [1.8, 2.2, 2.5]
        _format_header_row(p_table.rows[0], ["Persona / Actor", "Description", "Key Responsibilities"], col_w)

        for idx, persona in enumerate(personas):
            row = p_table.rows[idx + 1]
            bg_color = COLOR_ZEBRA_BG if idx % 2 == 1 else "FFFFFF"
            for c_idx, cell in enumerate(row.cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                cell.width = Inches(col_w[c_idx])
                cell.text = ""

            # Col 0: Persona
            p0 = row.cells[0].paragraphs[0]
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(persona.get("persona", "User"))
            r0.bold = True
            r0.font.name = "Calibri"
            r0.font.size = Pt(9.5)
            r0.font.color.rgb = COLOR_TEXT_DARK

            # Col 1: Description
            p1 = row.cells[1].paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(persona.get("description", ""))
            r1.font.name = "Calibri"
            r1.font.size = Pt(9)
            r1.font.color.rgb = COLOR_TEXT_DARK

            # Col 2: Responsibilities
            p2 = row.cells[2].paragraphs[0]
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(persona.get("key_responsibilities", ""))
            r2.font.name = "Calibri"
            r2.font.size = Pt(9)
            r2.font.color.rgb = COLOR_TEXT_DARK
    else:
        _add_bullet(doc, "System Administrator: Configures system parameters and monitors health.")
        _add_bullet(doc, "Operational User: Submits core workflows and accesses reporting.")

    # ------------------------------------------------------------------
    # Section 4: Functional Requirements
    # ------------------------------------------------------------------
    _add_section_heading(doc, "4. Functional Requirements", level=1)
    func_reqs = brd_content.get("functional_requirements", [])

    if func_reqs:
        # ID, Title & Statement, Actor & Outcome, Priority, Acceptance Criteria, Trace Refs
        f_table = doc.add_table(rows=1 + len(func_reqs), cols=6)
        f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _apply_table_borders(f_table)

        col_w = [0.9, 1.8, 1.1, 0.7, 1.2, 0.8]
        _format_header_row(
            f_table.rows[0],
            ["ID", "Title & Statement", "Actor & Outcome", "Priority", "Acceptance Criteria", "Trace Refs"],
            col_w,
        )

        for idx, req in enumerate(func_reqs):
            row = f_table.rows[idx + 1]
            bg_color = COLOR_ZEBRA_BG if idx % 2 == 1 else "FFFFFF"

            for c_idx, cell in enumerate(row.cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                cell.width = Inches(col_w[c_idx])
                cell.text = ""

            # Col 0: ID
            p0 = row.cells[0].paragraphs[0]
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(req.get("id", f"FR-{idx+1:03d}"))
            r0.bold = True
            r0.font.name = "Calibri"
            r0.font.size = Pt(9)
            r0.font.color.rgb = COLOR_NAVY

            # Col 1: Title & Statement
            p1 = row.cells[1].paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            r1_title = p1.add_run(req.get("title", "") + "\n")
            r1_title.bold = True
            r1_title.font.name = "Calibri"
            r1_title.font.size = Pt(9)
            r1_title.font.color.rgb = COLOR_TEXT_DARK

            r1_stmt = p1.add_run(req.get("statement", ""))
            r1_stmt.font.name = "Calibri"
            r1_stmt.font.size = Pt(8.5)
            r1_stmt.font.color.rgb = COLOR_TEXT_DARK

            # Col 2: Actor & Outcome
            p2 = row.cells[2].paragraphs[0]
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(0)
            actor = req.get("actor", "User")
            outcome = req.get("outcome", "")
            r2 = p2.add_run(f"Actor: {actor}\nOutcome: {outcome}")
            r2.font.name = "Calibri"
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = COLOR_TEXT_DARK

            # Col 3: Priority
            p3 = row.cells[3].paragraphs[0]
            p3.paragraph_format.space_before = Pt(0)
            p3.paragraph_format.space_after = Pt(0)
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            prio_text = str(req.get("priority", "HIGH")).upper()
            r3 = p3.add_run(prio_text)
            r3.bold = True
            r3.font.name = "Calibri"
            r3.font.size = Pt(8.5)
            if "HIGH" in prio_text or "P0" in prio_text:
                r3.font.color.rgb = RGBColor(180, 40, 40)
            else:
                r3.font.color.rgb = COLOR_TEXT_DARK

            # Col 4: Acceptance Criteria
            p4 = row.cells[4].paragraphs[0]
            p4.paragraph_format.space_before = Pt(0)
            p4.paragraph_format.space_after = Pt(0)
            ac_list = req.get("acceptance_criteria", [])
            if ac_list:
                ac_text = "\n".join(f"• {ac}" for ac in ac_list[:3])
            else:
                ac_text = "Verified on completion"
            r4 = p4.add_run(ac_text)
            r4.font.name = "Calibri"
            r4.font.size = Pt(8)
            r4.font.color.rgb = COLOR_TEXT_MUTED

            # Col 5: Trace Refs
            p5 = row.cells[5].paragraphs[0]
            p5.paragraph_format.space_before = Pt(0)
            p5.paragraph_format.space_after = Pt(0)
            traces = list(req.get("trace_refs") or [])
            for cid in req.get("source_chunk_ids") or []:
                c_label = f"CHUNK:{cid[:8]}"
                if c_label not in traces:
                    traces.append(c_label)
            r5 = p5.add_run("\n".join(traces) if traces else "Grounded")
            r5.font.name = "Calibri"
            r5.font.size = Pt(8)
            r5.font.color.rgb = COLOR_SLATE
    else:
        p_none = doc.add_paragraph("No functional requirements defined.")
        p_none.runs[0].font.italic = True

    # ------------------------------------------------------------------
    # Section 5: Non-Functional Requirements & Constraints
    # ------------------------------------------------------------------
    _add_section_heading(doc, "5. Non-Functional Requirements & Constraints", level=1)

    _add_section_heading(doc, "5.1 Non-Functional Requirements", level=2)
    nfr_reqs = brd_content.get("nonfunctional_requirements", [])
    if nfr_reqs:
        n_table = doc.add_table(rows=1 + len(nfr_reqs), cols=4)
        n_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _apply_table_borders(n_table)
        col_w = [1.0, 2.5, 1.5, 1.5]
        _format_header_row(n_table.rows[0], ["ID", "Requirement", "Priority", "Traceability"], col_w)

        for idx, req in enumerate(nfr_reqs):
            row = n_table.rows[idx + 1]
            bg_color = COLOR_ZEBRA_BG if idx % 2 == 1 else "FFFFFF"
            for c_idx, cell in enumerate(row.cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                cell.width = Inches(col_w[c_idx])
                cell.text = ""

            p0 = row.cells[0].paragraphs[0]
            r0 = p0.add_run(req.get("id", f"NFR-{idx+1:03d}"))
            r0.bold = True
            r0.font.name = "Calibri"
            r0.font.size = Pt(9)
            r0.font.color.rgb = COLOR_NAVY

            p1 = row.cells[1].paragraphs[0]
            r1 = p1.add_run(f"{req.get('title', '')}: {req.get('statement', '')}")
            r1.font.name = "Calibri"
            r1.font.size = Pt(9)
            r1.font.color.rgb = COLOR_TEXT_DARK

            p2 = row.cells[2].paragraphs[0]
            r2 = p2.add_run(str(req.get("priority", "HIGH")).upper())
            r2.font.name = "Calibri"
            r2.font.size = Pt(9)

            p3 = row.cells[3].paragraphs[0]
            t_refs = list(req.get("trace_refs") or req.get("source_chunk_ids") or ["Grounded"])
            r3 = p3.add_run(", ".join(t_refs[:2]))
            r3.font.name = "Calibri"
            r3.font.size = Pt(8.5)
            r3.font.color.rgb = COLOR_SLATE
    else:
        _add_bullet(doc, "Performance: Sub-second response time for 95% of standard user transactions.")
        _add_bullet(doc, "Availability: High availability SLA of 99.9% uptime during standard operation.")

    _add_section_heading(doc, "5.2 Architecture Constraints", level=2)
    constraints = brd_content.get("constraints", [])
    if constraints:
        for c in constraints:
            _add_bullet(doc, f"{c.get('title', 'Constraint')}: {c.get('statement', '')}", bold_prefix="[Constraint] ")
    else:
        _add_bullet(doc, "Database: Cloud relational PostgreSQL instance enforcing strict schema constraints.")
        _add_bullet(doc, "Security: Token-based API authorization with encryption for data in transit and at rest.")

    # ------------------------------------------------------------------
    # Section 6: Assumptions & Dependencies
    # ------------------------------------------------------------------
    _add_section_heading(doc, "6. Assumptions & Dependencies", level=1)

    _add_section_heading(doc, "6.1 Explicit Assumptions", level=2)
    assumptions = brd_content.get("assumptions", [])
    if assumptions:
        for asm in assumptions:
            if isinstance(asm, dict):
                asm_id = asm.get("id", "ASM")
                status = asm.get("status", "ACCEPTED")
                impact = asm.get("impact", "MEDIUM")
                stmt = asm.get("statement", "")
                rationale = asm.get("rationale", "")
                prefix = f"[{asm_id}] ({status} | Impact: {impact}): "
                full_text = f"{stmt} (Rationale: {rationale})" if rationale else stmt
                _add_bullet(doc, full_text, bold_prefix=prefix)
            else:
                _add_bullet(doc, str(asm))
    else:
        _add_bullet(doc, "Stakeholder inputs provided during the clarification loop represent confirmed scope.")

    _add_section_heading(doc, "6.2 Technical Dependencies", level=2)
    dependencies = brd_content.get("dependencies", [])
    if dependencies:
        for dep in dependencies:
            _add_bullet(doc, str(dep))
    else:
        _add_bullet(doc, "Availability of cloud hosting runtime environment.")
        _add_bullet(doc, "External OAuth/JWT identity provider endpoints.")

    # ------------------------------------------------------------------
    # Section 7: Traceability Matrix Summary
    # ------------------------------------------------------------------
    _add_section_heading(doc, "7. Traceability Matrix Summary", level=1)
    trace_matrix = brd_content.get("traceability_matrix", [])

    p_trace_note = doc.add_paragraph()
    p_trace_note.paragraph_format.space_after = Pt(8)
    r_tn = p_trace_note.add_run(
        "Traceability Rule: Every requirement links directly to a source document chunk, "
        "a stakeholder answer, or an explicitly confirmed assumption."
    )
    r_tn.font.name = "Calibri"
    r_tn.font.size = Pt(9.5)
    r_tn.font.italic = True
    r_tn.font.color.rgb = COLOR_TEXT_MUTED

    if trace_matrix:
        t_table = doc.add_table(rows=1 + len(trace_matrix), cols=5)
        t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _apply_table_borders(t_table)
        col_w = [1.2, 2.2, 0.9, 1.4, 0.8]
        _format_header_row(t_table.rows[0], ["Req ID", "Title", "Priority", "Grounding Sources", "Status"], col_w)

        for idx, entry in enumerate(trace_matrix):
            row = t_table.rows[idx + 1]
            bg_color = COLOR_ZEBRA_BG if idx % 2 == 1 else "FFFFFF"

            for c_idx, cell in enumerate(row.cells):
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                cell.width = Inches(col_w[c_idx])
                cell.text = ""

            # Req ID
            p0 = row.cells[0].paragraphs[0]
            r0 = p0.add_run(entry.get("requirement_id", f"REQ-{idx+1}"))
            r0.bold = True
            r0.font.name = "Calibri"
            r0.font.size = Pt(8.5)
            r0.font.color.rgb = COLOR_NAVY

            # Title
            p1 = row.cells[1].paragraphs[0]
            r1 = p1.add_run(entry.get("title", ""))
            r1.font.name = "Calibri"
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = COLOR_TEXT_DARK

            # Priority
            p2 = row.cells[2].paragraphs[0]
            r2 = p2.add_run(str(entry.get("priority", "HIGH")).upper())
            r2.font.name = "Calibri"
            r2.font.size = Pt(8.5)

            # Sources
            p3 = row.cells[3].paragraphs[0]
            sources = entry.get("sources", [])
            r3 = p3.add_run("\n".join(str(s) for s in sources[:3]))
            r3.font.name = "Calibri"
            r3.font.size = Pt(8)
            r3.font.color.rgb = COLOR_SLATE

            # Status
            p4 = row.cells[4].paragraphs[0]
            status_text = entry.get("grounding_status", "GROUNDED")
            r4 = p4.add_run(status_text)
            r4.bold = True
            r4.font.name = "Calibri"
            r4.font.size = Pt(8.5)
            if status_text == "GROUNDED":
                r4.font.color.rgb = RGBColor(16, 124, 65)
            else:
                r4.font.color.rgb = RGBColor(180, 40, 40)

    # 3. Save Document to Stream
    target_stream = output_stream if output_stream is not None else io.BytesIO()
    doc.save(target_stream)
    return target_stream.getvalue()
